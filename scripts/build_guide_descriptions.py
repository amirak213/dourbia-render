import re, sys, unicodedata
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "AgentHistorique" / "data" / "raw" / "excel" / "Monuments.xlsx"
def _default_out() -> Path:
    for base in (Path(r"C:\Users\Hp23\dourbia_dumps"), Path(r"E:\dourbia_dumps")):
        d = base / "utf8"
        if base.is_dir():
            d.mkdir(parents=True, exist_ok=True)
            return d / "guide_descriptions.sql"
    out = ROOT / "dourbia_dumps" / "utf8" / "guide_descriptions.sql"
    out.parent.mkdir(parents=True, exist_ok=True)
    return out

# id base -> (ID_monument Excel, nom normalise attendu), d'apres l'audit R2
MAP = {
    1: ("1", "theatre"),
    2: ("2", "parcdesvillasromaines"),
    12: ("3", "parcdesthermesdantonin"),
    23: ("4", "quartiermagon"),
    24: ("5", "tophet"),
    25: ("6", "portspuniques"),
    43: ("9", "beitelhikma"),
    44: ("10", "museeoceanographiquedarelhout"),
}

def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "").casefold())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", s)

def sid(v):
    s = "" if v is None else str(v).strip()
    return s[:-2] if s.endswith(".0") else s

OUT = _default_out()
ws = load_workbook(XLSX, read_only=True, data_only=True)["Monuments"]
rows = list(ws.iter_rows(values_only=True))
ix = {str(h).strip(): i for i, h in enumerate(rows[0]) if h}
need = ["ID_monument", "nom_monument_FR", "description_FR", "description_EN", "description_AR"]
missing = [c for c in need if c not in ix]
if missing:
    sys.exit(f"Colonnes absentes: {missing}")

lines = ["BEGIN;"]
for base_id, (xid, expected) in MAP.items():
    hits = [r for r in rows[1:] if sid(r[ix["ID_monument"]]) == xid]
    if len(hits) != 1:
        sys.exit(f"ID Excel {xid}: {len(hits)} lignes (attendu 1)")
    r = hits[0]
    if norm(r[ix["nom_monument_FR"]]) != expected:
        sys.exit(f"ID Excel {xid}: nom inattendu {r[ix['nom_monument_FR']]!r}")
    for lang in ("FR", "EN", "AR"):
        text = str(r[ix[f"description_{lang}"]] or "").strip()
        if len(text) < 20:
            print(f"base {base_id} {lang}: vide, ignore")
            continue
        if "$g$" in text:
            sys.exit(f"base {base_id} {lang}: contient $g$")
        col = f"description_{lang.lower()}"
        lines.append(
            f"UPDATE monuments SET {col} = $g${text}$g$ "
            f"WHERE id = {base_id} AND ({col} IS NULL OR btrim({col}) = '');"
        )
        print(f"base {base_id} {lang}: {len(text)} caracteres")
lines.append("COMMIT;")
OUT.write_bytes(("\n".join(lines) + "\n").encode("utf-8"))
print("Ecrit:", OUT)

